"""Збирає votes.js для сайту з реєстру власників (registry_*.xlsx) і експорту ДАХ.

Голос = власник зареєстрований у ДАХ **і** має підтверджене право власності
(колонка F «Статус» = «Право власності»). Рахуються квартири й нежитлові приміщення
(база ОСББ), паркомісця — ні.

На сайт потрапляють ЛИШЕ агрегати: скільки голосів по кожній секції і загальна площа.
Номери квартир, поверхи, імена й телефони не публікуються (сайт підсвічує вікна випадково).

Запуск:
  python3 tools/build_votes.py registry.xlsx dah_export.xlsx   # «в ДАХ» — з експорту «Реєстр приміщень» ДАХ
  python3 tools/build_votes.py registry.xlsx                   # «в ДАХ» — з колонки G реєстру (і E нежитлових)
"""
import json
import re
import sys
from datetime import datetime
from statistics import median

import openpyxl

APTS, NONRES = "Реєстр мешканців", "Нежитлові приміщення"
IN_DAH, OWNER = "✅ У ДАХ", "Право власності"
BASE_AREA = 50258.3  # квартири + нежитлові


def section_starts():
    """Перші квартири секцій беремо з data.js (`first: N`), щоб не дублювати розкладку."""
    with open("data.js", encoding="utf-8") as f:
        return sorted(int(n) for n in re.findall(r"first:\s*(\d+)", f.read()))


def area(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", ".").strip())
        except ValueError:
            return None
    return None


def key(v):
    """1.0 -> "1", "337-1" -> "337-1", 1004.0 -> "1004", інше -> None."""
    if isinstance(v, (int, float)):
        return str(int(v))
    if isinstance(v, str) and re.fullmatch(r"\d+(-\d+)?", v.strip()):
        return v.strip()
    return None


def dah_owners(path):
    """Ключі квартир і нежитлових, у яких в експорті ДАХ заповнений власник."""
    apts, nonres = set(), set()
    for r in openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0].iter_rows(values_only=True):
        if not r or not r[0] or len(r) < 9 or not r[8]:
            continue
        name = str(r[0])
        if m := re.search(r"Квартира\s*([\d-]+)", name):
            apts.add(m[1])
        elif m := re.search(r"Нежитлове приміщення\s*(\d+)", name):
            nonres.add(m[1])
    return apts, nonres


def main(reg_path, dah_path=None):
    wb = openpyxl.load_workbook(reg_path, data_only=True, read_only=True)
    apts = [r for r in wb[APTS].iter_rows(values_only=True) if r and key(r[0])]
    nonres = [r for r in wb[NONRES].iter_rows(values_only=True) if r and key(r[0])] if NONRES in wb.sheetnames else []
    dah_a, dah_n = dah_owners(dah_path) if dah_path else (None, None)

    def votes(r, dah_set, dah_col):
        in_dah = key(r[0]) in dah_set if dah_set is not None else (len(r) > dah_col and r[dah_col] == IN_DAH)
        return in_dah and r[3 if dah_col == 4 else 5] == OWNER  # статус: F у квартир, D у нежитлових

    fill = median(a for a in (area(r[1]) for r in apts) if a)  # якщо в реєстрі немає площі
    starts = section_starts()
    sections = [{"total": 0, "inDah": 0} for _ in starts]
    vote_area, estimated = 0.0, 0
    for r in apts:
        n = int(key(r[0]).split("-")[0])
        sec = sections[max(i for i, f in enumerate(starts) if n >= f)]
        sec["total"] += 1
        if votes(r, dah_a, 6):
            sec["inDah"] += 1
            a = area(r[1])
            if a is None:
                a, estimated = fill, estimated + 1
            vote_area += a
    nonres_votes = [r for r in nonres if votes(r, dah_n, 4)]
    nonres_area = sum(area(r[1]) or 0 for r in nonres_votes)

    out = {
        "updatedAt": datetime.now().strftime("%Y-%m-%dT%H:%M"),
        "baseArea": BASE_AREA,
        "dahArea": round(vote_area + nonres_area, 1),
        "areaEstimated": estimated,
        "nonres": len(nonres_votes),
        "sections": sections,
    }
    with open("votes.js", "w", encoding="utf-8") as f:
        f.write("// Згенеровано tools/build_votes.py — не редагуйте вручну.\n")
        f.write("window.VOTES = " + json.dumps(out, ensure_ascii=False, separators=(",", ":")) + ";\n")
    n = sum(x["inDah"] for x in sections)
    print(f"Голоси (в ДАХ + право власності): квартир {n} ({vote_area:.1f} м², по секціях {[x['inDah'] for x in sections]}), "
          f"нежитлових {len(nonres_votes)} ({nonres_area:.1f} м²) → {out['dahArea']} м² = {out['dahArea'] / BASE_AREA:.1%}; "
          f"площу оцінено для {estimated} кв.")


if __name__ == "__main__":
    main(*sys.argv[1:3])
